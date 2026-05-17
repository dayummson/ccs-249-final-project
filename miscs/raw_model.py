import docx
from transformers import pipeline
import re
import docx
from pypdf import PdfReader

reader = PdfReader("./samples/final-project.pdf")

MODEL = "distilbert-base-uncased"


def extract_smart_blocks(file_path):
    doc = docx.Document(file_path)
    # Filter out empty lines and headers that are too short to be tasks
    blocks = [p.text.strip() for p in doc.paragraphs if len(p.text.strip()) > 15]
    return blocks


pipe = pipeline("text-classification", model=MODEL)


# def extract_text_from_docx(file):
#     doc = docx.Document(file)
#     return " ".join([para.text for para in doc.paragraphs])


raw_text = extract_smart_blocks("./samples/unit-4-activity.docx")

print(raw_text)

id_2_label = {
    "LABEL_0": "ADMIN_TRAP",
    "LABEL_1": "IGNORE",
    "LABEL_2": "OPTIONAL_BONUS",
    "LABEL_3": "PRE_REQUISITE",
    "LABEL_4": "TECHNICAL_TASK",
    "LABEL_5": "WEIGHTED_PRIORITY",
}


def process_to_briefing(blocks):
    results = pipe(blocks)

    print("[RESULTS]: ", results)

    briefing = {
        "ADMIN_TRAP": [],
        "TECHNICAL_TASK": [],
        "WEIGHTED_PRIORITY": [],
        "OPTIONAL_BONUS": [],
    }

    for text, res in zip(blocks, results):
        label = id_2_label.get(res["label"], "UNKNOWN")

        # HEURISTIC OVERRIDE: Rubrics are predictable. Let's use that.
        low_text = text.lower()
        if "points" in low_text or "pts" in low_text:
            if "bonus" in low_text:
                label = "OPTIONAL_BONUS"
            else:
                label = "WEIGHTED_PRIORITY"
        elif any(
            verb in low_text
            for verb in ["implement", "extract", "modify", "write", "code"]
        ):
            label = "TECHNICAL_TASK"

        if label in briefing:
            briefing[label].append(text)

    return briefing


def write_to_txt(briefing, filename="briefing.txt"):
    with open(filename, "w") as f:
        for category, items in briefing.items():
            f.write(f"{category}:\n")
            for item in items:
                f.write(f" - {item}\n")
            f.write("\n")


OUTPUT_FILE = "raw_model_output.txt"

briefing = process_to_briefing(raw_text)
write_to_txt(briefing, filename=OUTPUT_FILE)
print(f"Briefing generated and saved to {OUTPUT_FILE}")
